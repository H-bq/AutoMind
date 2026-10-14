from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
OUT_PATH = OUT_DIR / "AutoMind-作品设计文档.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY_FILL = "F2F4F7"
BORDER = "D9E2F3"


def set_east_asian_font(run, font_name="微软雅黑"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)


def set_table_width(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(int(width * 1440)))


def style_paragraph(paragraph, size=11, color=None, bold=False, after=6, before=0, line=1.1):
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.line_spacing = line
    for run in paragraph.runs:
        set_east_asian_font(run)
        run.font.size = Pt(size)
        run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color


def add_para(doc, text="", size=11, bold=False, color=None, after=6, before=0, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_east_asian_font(r)
    r.font.size = Pt(size)
    r.font.bold = bold
    if color is not None:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.1
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_east_asian_font(r)
    if level == 1:
        r.font.size = Pt(16)
        r.font.color.rgb = BLUE
        r.font.bold = True
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        r.font.size = Pt(13)
        r.font.color.rgb = BLUE
        r.font.bold = True
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    else:
        r.font.size = Pt(12)
        r.font.color.rgb = DARK_BLUE
        r.font.bold = True
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style=None)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run("· " + text)
    set_east_asian_font(r)
    r.font.size = Pt(11)
    return p


def add_numbered(doc, number, text):
    p = doc.add_paragraph(style=None)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(f"{number}. {text}")
    set_east_asian_font(r)
    r.font.size = Pt(11)
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    set_table_width(table, widths)
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        r = p.add_run(text)
        set_east_asian_font(r)
        r.font.size = Pt(10.5)
        r.font.bold = True
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(hdr[i], GRAY_FILL)
        set_cell_margins(hdr[i])
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(text)
            set_east_asian_font(r)
            r.font.size = Pt(10)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
    set_table_width(table, widths)
    add_para(doc, "", after=2)
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(11)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("AutoMind 作品设计文档")
    set_east_asian_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(100, 100, 100)


def build():
    OUT_DIR.mkdir(exist_ok=True)
    doc = Document()
    configure_document(doc)

    title = add_para(doc, "AutoMind：面向智网意图场景的跨域分布式多智能体调度系统", size=20, bold=True, color=DARK_BLUE, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    subtitle = add_para(doc, "作品设计文档", size=16, bold=True, color=BLUE, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "匿名说明：本文档遵循模板要求，不包含学校、团队及队员身份等相关信息。", size=10.5, color=RGBColor(90, 90, 90), after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "所在赛道与赛项：待确认（请在 A / A-ST / A-ICV / B-EP1 中选择后替换）。", size=11, bold=True, after=12)

    add_heading(doc, "一、目标问题与意义价值", 1)
    add_para(doc, "AutoMind 面向“智网意图”场景，关注业务用户以自然语言提出网络质量保障需求后，如何将模糊业务语言转化为可校验、可审计、可执行的网络策略执行闭环。系统当前不直接接入真实物理网络设备或生产流量，而是在软件模拟环境中完成意图解析、多智能体协同、MCP 工具调用、SLA 遥测评估、异常注入与自愈恢复演示。")
    add_bullet(doc, "应用领域：面向企业专线、跨域园区网、教育直播、金融结算、制造视觉巡检等需要确定性网络服务质量的场景。")
    add_bullet(doc, "关注问题：传统网络配置依赖人工理解业务目标、手工检索拓扑与资源状态，策略下发与 SLA 验证之间缺少连续审计链路。")
    add_bullet(doc, "实现目标：通过 Supervisor、Planner、Executor、Monitor 多智能体完成 Plan-Execute-Verify 闭环，使自然语言意图转化为结构化计划、模拟策略执行、SLA 验证和自愈重规划。")
    add_bullet(doc, "应用价值：降低业务侧表达网络需求的门槛，提高跨域网络调度的可解释性、可追踪性和演示验证效率，为后续接入真实 SDN 控制器或网络仿真平台预留工程边界。")

    add_heading(doc, "二、设计思路与方案", 1)
    add_heading(doc, "2.1 总体技术路线", 2)
    add_para(doc, "系统采用 Maven 多模块单体与独立模拟器进程结合的方案：AutoMind-service 作为智能编排与业务 API 中心，AutoMind-common 承载通用异常、上下文、响应模型和工具能力，AutoMind-simulator 作为独立遥测与 SLA 数据生成进程。主服务端口为 9090，模拟器端口为 9091，二者通过内部令牌保护的 HTTP 接口传递实时 SLA 事件。")
    add_numbered(doc, 1, "意图入口：用户通过 /intent-request/submit 提交自然语言需求，系统记录租户、用户、角色和 traceId。")
    add_numbered(doc, 2, "智能编排：Supervisor 解析意图并输出置信度；Planner 生成路径与资源计划；Executor 通过 MCP 调用模拟控制能力；Monitor 读取遥测并形成 SLA 结论。")
    add_numbered(doc, 3, "管控闭环：低置信度、跨域降级、高风险 MCP 操作或高风险计划进入审批队列，审批通过后自动恢复编排。")
    add_numbered(doc, 4, "运行验证：模拟器按周期生成正常遥测、按需注入异常，服务端通过 SLA 记录、实时事件流和自愈记录呈现运行效果。")

    add_heading(doc, "2.2 架构组成", 2)
    add_table(
        doc,
        ["层次", "核心组件", "设计说明"],
        [
            ["用户与管理入口", "用户认证、意图提交、审批队列、拓扑与 SLA 查询接口", "区分业务用户、管理员、运维和审计角色，接口通过租户上下文进行数据隔离。"],
            ["Intelligence Plane", "Supervisor、Planner、Executor、Monitor、PEV 状态机", "围绕意图生命周期组织 Agent 任务，记录输入、输出、置信度、错误码和执行时间。"],
            ["工具与控制层", "MCP Client、Mock MCP Server、get_topo / set_qos / block_flow", "使用标准化工具注册、角色校验、超时分类、重试熔断和调用日志模拟 SDN 控制器能力。"],
            ["跨域协同层", "A2A Message、Sender/Receiver、FallbackGuard", "用于跨网络域交换链路状态；通信失败时进入本域自治降级，并保留审计消息。"],
            ["监控与自愈层", "SLA Evaluator、HealingEngine、SSE 实时推送、Simulator", "检测 SLA 违约、MCP 超时、A2A 降级等事件，触发有限重规划或人工接管。"],
            ["数据与审计层", "MySQL、Redis、mcp_call_log、agent_task、approval_record、healing_record", "将意图、任务、审批、工具调用、SLA 和遥测数据用 traceId 与 intentId 关联。"],
        ],
        [1.3, 1.8, 3.2],
    )

    add_heading(doc, "2.3 关键流程", 2)
    add_para(doc, "核心业务流程可概括为“自然语言意图 -> 结构化解析 -> 路径规划 -> 工具执行 -> SLA 验证 -> 自愈/完成”。当 Supervisor 发现字段缺失或置信度低于阈值，系统进入补充信息流程；当 Planner 产生高风险方案或跨域降级结果，系统进入审批流程；当 Monitor 判定 SLA 违约，HealingEngine 根据触发源执行有限次数重规划，超过最大次数后进入人工接管。")
    add_table(
        doc,
        ["流程阶段", "输入", "处理动作", "输出/状态"],
        [
            ["意图提交", "自然语言需求、用户/租户上下文", "写入 intent_request，异步启动编排", "SUBMITTED / traceId"],
            ["Supervisor", "原始文本", "抽取带宽、时延、时间、链路等字段并判断置信度", "ParsedIntent / SUPERVISED"],
            ["Planner", "结构化意图、拓扑、资源状态", "生成主备路径与资源预留计划，识别风险", "PlanDecision / AWAIT_APPROVAL 或继续执行"],
            ["Executor", "计划结果", "调用 MCP 工具模拟查询拓扑、设置 QoS、阻断流量", "ExecutionResult / MCP 审计日志"],
            ["Monitor", "执行结果、遥测数据", "评估 LATENCY、JITTER、LOSS、BANDWIDTH 指标", "SlaVerdict / COMPLETED 或 HEALING"],
            ["Healing", "SLA 违约、MCP 超时、A2A 降级事件", "有限次数重规划，本域降级或人工接管", "healing_record / HANDOFF_HUMAN"],
        ],
        [1.05, 1.45, 2.75, 1.25],
    )

    add_heading(doc, "三、方案实现", 1)
    add_heading(doc, "3.1 软件技术与集成方法", 2)
    add_table(
        doc,
        ["类别", "选型/实现", "作用"],
        [
            ["后端框架", "Java 17、Spring Boot 3.4.3、Spring AI 1.1.0、Spring AI Alibaba 1.1.0.0-RC2", "支撑 REST API、多智能体编排、DashScope/Qwen-plus 兼容调用和结构化输出。"],
            ["数据访问", "MyBatis-Plus 3.5.7、MySQL 8.0、H2 测试库", "管理网络域、节点、链路、意图、审计、SLA 和遥测数据。"],
            ["可靠性", "Resilience4j Retry/CircuitBreaker、异步线程池、traceId MDC", "处理 AI/MCP 调用超时、有限重试、熔断和链路追踪。"],
            ["模拟与验证", "AutoMind-simulator、周期调度、异常注入、SSE 实时事件", "提供不依赖真实设备的可演示运行环境。"],
            ["接口文档", "SpringDoc OpenAPI", "通过 /swagger-ui.html 和 /v3/api-docs 查看接口契约。"],
        ],
        [1.15, 2.35, 2.8],
    )

    add_heading(doc, "3.2 功能模块实现", 2)
    add_table(
        doc,
        ["模块", "当前实现内容", "对应代码事实"],
        [
            ["意图与用户", "登录、注册、意图提交、意图状态查询、全链路追踪", "/user、/intent-request、/intent-request/{id}/trace、/trace/{traceId}"],
            ["多智能体编排", "Supervisor/Planner/Executor/Monitor 任务记录、PEV 状态推进、审批恢复", "IntentOrchestrationService、BaseAgent、ApprovalResumeListener"],
            ["审批与审计", "待审批列表、审批历史、批准、拒绝、拒绝并要求修改", "/approvals/pending、/approvals/history、/approve、/reject-with-modification"],
            ["MCP 工具调用", "工具注册、角色校验、超时重试、错误分类、调用日志", "McpClient、McpToolRegistry、MockMcpServerController、mcp_call_log"],
            ["A2A 协同", "跨域消息、消息审计、通信降级事件", "A2aMessageSender、A2aMessageReceiver、A2aFallbackGuard、/a2a-messages"],
            ["SLA 与实时监控", "SLA 仪表盘、单次评估、历史记录、SSE 订阅", "/sla/dashboard、/sla/evaluate、/sla/records、/realtime/sla/stream"],
            ["模拟器", "正常遥测周期生成、SLA 周期评估、手动异常注入、过期数据清理", "/simulator/status、/start、/stop、/anomaly/generate、/cleanup/run"],
        ],
        [1.25, 2.7, 2.45],
    )

    add_heading(doc, "3.3 数据模型与审计链路", 2)
    add_para(doc, "数据设计围绕 intentId、agentTaskId 和 traceId 建立可追踪关系。核心表包括 user、person、intent_request、agent_task、approval_record、mcp_call_log、a2a_message、audit_log、network_domain、network_node、network_link、network_link_resource、network_resource_reservation、telemetry_data、sla_record、healing_record 等。为了支撑实时 SLA 查询，系统补充了 telemetry_data(time)、telemetry_data(tenant_id, link_id, time)、sla_record(evaluated_at)、sla_record(tenant_id, link_id, evaluated_at) 等索引。")
    add_bullet(doc, "租户隔离：UserContext 在请求与 Agent 执行链中携带 tenantId、role、userId，核心查询按租户过滤，admin 租户可查看全量。")
    add_bullet(doc, "审计留痕：MCP 调用记录包含工具名、版本、参数、响应码、错误类型、耗时、尝试次数和高风险标记。")
    add_bullet(doc, "实时关联：模拟器生成 telemetry_data 和 sla_record 后，通过 /internal/sla-events 推送事件，前端或管理端可订阅 /realtime/sla/stream。")

    add_heading(doc, "四、运行结果/应用效果", 1)
    add_para(doc, "当前作品已形成可运行的软件演示闭环：主服务负责意图、Agent、审批、MCP、A2A、SLA 与自愈；模拟器负责生成遥测和异常注入。默认配置下主服务运行在 9090 端口，模拟器运行在 9091 端口；模拟器每 5 秒生成一轮正常遥测，每 8 秒评估 SLA，每 30 秒执行过期数据清理，默认评估窗口为 5 分钟。")
    add_table(
        doc,
        ["验证项", "结果", "说明"],
        [
            ["服务端单元/集成测试", "12/13 通过", "mvn test 中 ApprovalAspect、StageFourEndToEnd、StructuredOutput、SSE 控制器等测试通过；InternalSlaEventControllerTest 有 1 个用例因测试构造未注入 IPersonService 触发 NPE，属于待修复测试问题。"],
            ["模拟器测试", "5/5 通过", "mvn test -pl AutoMind-simulator -am 通过，覆盖异常注入、清理、SLA 评估、遥测样本生成。"],
            ["SLA 实时能力", "已实现", "内部事件接口 /internal/sla-events 校验令牌后发布 SSE 事件，订阅入口为 /realtime/sla/stream?token=...。"],
            ["异常演示能力", "已实现", "通过 /simulator/anomaly/generate 手动生成异常突发样本，并触发 SLA 记录生成。"],
            ["可审计链路", "已实现", "意图、Agent 任务、MCP 调用、审批、A2A 消息、SLA 记录和自愈记录可通过 traceId / intentId 关联。"],
        ],
        [1.35, 1.15, 3.85],
    )
    add_para(doc, "应用效果上，系统能够把“保障某业务链路低时延/低丢包/带宽质量”的自然语言需求拆解为结构化意图，并通过多智能体任务、MCP 工具调用和 SLA 评估记录呈现全过程。对于异常链路，系统可通过模拟器注入高时延、高丢包或高带宽占用样本，在 SLA 评估中形成违约记录，并触发重规划或人工接管流程。")

    add_heading(doc, "五、创新与特色", 1)
    add_numbered(doc, 1, "面向智网意图的 PEV 多智能体闭环：将自然语言意图解析、路径规划、策略执行和 SLA 验证拆成四类 Agent，既保留大模型理解能力，又通过任务表和状态机控制执行边界。")
    add_numbered(doc, 2, "MCP 工具调用与网络控制审计结合：使用工具注册、角色校验、高风险标记、超时重试和调用日志，把模拟 SDN 控制动作纳入可追踪、可解释的工程链路。")
    add_numbered(doc, 3, "A2A 跨域协同与降级自治：通过跨域消息交换链路状态，当跨域通信失败时进入本域自治模式，并通过降级事件触发后续自愈流程。")
    add_numbered(doc, 4, "独立模拟器支撑可演示验证：将遥测生成、SLA 周期评估、异常注入和数据清理放入独立进程，避免主服务承载演示数据压力，也便于后续替换为真实仿真环境。")
    add_numbered(doc, 5, "全链路安全与可观测设计：从用户认证、租户隔离、审批队列、审计日志到 traceId 串联，形成从提交、决策、执行、验证到恢复的证据链。")

    add_heading(doc, "六、后续完善方向", 1)
    add_para(doc, "作品当前以软件模拟和演示验证为主，后续可沿三条路径增强：第一，修复 InternalSlaEventControllerTest 的测试构造问题并补齐核心 Service 单元测试；第二，将 staging/prod 配置、密钥外部注入和 CI/CD 流水线纳入工程化交付；第三，将 Mock MCP Server 逐步替换为真实 SDN 控制器或网络仿真平台适配器，验证跨域调度在更接近真实网络环境中的效果。")

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
